import os
import io
import json
import csv
import logging
from typing import Optional, Dict, Any

from src.tools.registry import register_tool

logger = logging.getLogger(__name__)

# ==========================================
# Universal Document & File Parser
# ==========================================

@register_tool(
    name="read_document_file",
    description="استخراج و خواندن متن کامل از فایل‌های متنی، اسناد PDF، فایل‌های Word (DOCX) و اکسل (XLSX)",
    category="files"
)
def read_document_file(file_path: str, max_chars: int = 4000) -> str:
    """
    :param file_path: مسیر فایل ذخیره‌شده
    :param max_chars: حداکثر تعداد کاراکترهای استخراجی
    """
    clean_path = file_path.strip()
    # Security Sandbox Protection: Never expose bot source code or sensitive tokens/configs
    norm_p = os.path.abspath(clean_path)
    forbidden_sources = [
        "bot.py", ".env", "config.py", "database.py", "ai_service.py",
        "/etc/shadow", "/etc/passwd", "token", "secret", "private_key", ".git"
    ]
    if any(fb in norm_p.lower() for fb in forbidden_sources) or norm_p.endswith((".py", ".env", ".pem", ".key")):
        return "⛔ دسترسی غیرمجاز: خواندن سورس‌کدهای اصلی ربات، کلیدها یا پیکربندی‌های محرمانه سیستم اکیداً مسدود است."

    if not os.path.exists(clean_path):
        return f"فایل در مسیر '{file_path}' یافت نشد."

    ext = os.path.splitext(clean_path)[1].lower()
    extracted_text = ""

    try:
        if ext == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(clean_path)
            pages_text = [page.extract_text() or "" for page in reader.pages[:15]]
            extracted_text = "\n\n".join(pages_text)

        elif ext == ".docx":
            import docx
            doc = docx.Document(clean_path)
            extracted_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

        elif ext in [".xlsx", ".xls"]:
            import openpyxl
            try:
                wb = openpyxl.load_workbook(clean_path, data_only=True)
            except Exception:
                return "فایل اکسل قدیمی (.xls) پشتیبانی نمی‌شود؛ لطفاً آن را به .xlsx تبدیل کنید."
            sheet = wb.active
            rows_data = []
            for row in sheet.iter_rows(values_only=True):
                r_str = " | ".join(str(cell) for cell in row if cell is not None)
                if r_str:
                    rows_data.append(r_str)
            extracted_text = "\n".join(rows_data[:100])

        elif ext == ".zip":
            import zipfile
            try:
                with zipfile.ZipFile(clean_path, "r") as zf:
                    names = zf.namelist()[:50]
                    extracted_text = "\n".join(f"{'📁' if not n.endswith('/') else '📂'} {n} ({zf.getinfo(n).file_size:,} bytes)" if not n.endswith("/") else f"📂 {n}" for n in names)
                    if len(zf.namelist()) > 50:
                        extracted_text += f"\n... و {len(zf.namelist()) - 50} فایل دیگر"
            except zipfile.BadZipFile:
                return "فایل ZIP خراب است و باز نمی‌شود."

        else:
            with open(clean_path, "r", encoding="utf-8", errors="replace") as f:
                extracted_text = f.read(max_chars)

        if not extracted_text.strip():
            return "محتوای متنی قابل استخراجی در فایل یافت نشد."

        if len(extracted_text) > max_chars:
            extracted_text = extracted_text[:max_chars] + "\n\n...[ادامه محتوای فایل به دلیل محدودیت خلاصه گردید]"

        return f"📄 *محتوای استخراج‌شده از سند ({os.path.basename(clean_path)})*:\n\n{extracted_text}"

    except Exception as e:
        return f"خطا در خواندن فایل: {str(e)}"

# ==========================================
# Universal File Generator & Telegram Uploader
# ==========================================

@register_tool(
    name="create_and_upload_file",
    description="تولید خودکار و هوشمند فایل استاندارد با فرمت‌های مختلف (مانند PDF, Word/DOCX, Excel/XLSX, CSV, JSON, TXT) و ارسال مستقیم آن در تلگرام. توجه: درون پارامتر content باید خود متن و مقاله‌ای که باید درون فایل باشد قرار گیرد، نه کد برای ساختن آن.",
    category="files"
)
async def create_and_upload_file(
    filename: str,
    content: str,
    caption: str = ""
) -> Dict[str, Any]:
    """
    :param filename: نام فایل به همراه پسوند (مثال: report.pdf, data.xlsx, script.py, notes.txt, document.docx)
    :param content: متن، جدول، کد یا محتوای مورد نظر برای قرار گرفتن درون فایل
    :param caption: کپشن یا توضیح اختیاری فایل برای ارسال در تلگرام
    """
    clean_name = (filename or "").strip().replace("/", "_").replace(chr(92), "_")
    if not clean_name:
        clean_name = "file.txt"
    if len(content or "") > 400000:
        return {
            "type": "error",
            "status": "error",
            "message": "❌ حجم متن ورودی بیش از حد مجاز است (سقف ۴۰۰ هزار نویسه). لطفاً متن را کوتاه‌تر کنید."
        }

    base_name, ext = os.path.splitext(clean_name)
    ext = ext.lower()
    if not ext:
        ext = ".txt"
        clean_name += ".txt"

    file_bytes = io.BytesIO()

    try:
        if ext == ".pdf":
            # High-Fidelity Persian & Multilingual Multi-Page PDF Document Engine
            import textwrap
            import re
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont

            # Check and register Persian TrueType font (Vazirmatn)
            has_vazir = False
            vazir_candidates = [
                os.path.join(os.getcwd(), "assets", "fonts", "Vazirmatn-Regular.ttf"),
                "/home/dsh/workspace/prometheuspro/assets/fonts/Vazirmatn-Regular.ttf"
            ]
            for fpath in vazir_candidates:
                if os.path.exists(fpath):
                    try:
                        pdfmetrics.registerFont(TTFont("Vazirmatn", fpath))
                        has_vazir = True
                        break
                    except Exception:
                        pass

            def _is_rtl(text_s: str) -> bool:
                return any('\u0600' <= ch <= '\u06FF' or '\uFB50' <= ch <= '\uFDFF' or '\uFE70' <= ch <= '\uFEFF' for ch in text_s)

            def _shape_persian(text_s: str) -> str:
                if not text_s:
                    return ""
                try:
                    import arabic_reshaper
                    from bidi.algorithm import get_display
                    reshaped = arabic_reshaper.reshape(text_s)
                    return get_display(reshaped)
                except Exception:
                    return text_s

            c = canvas.Canvas(file_bytes, pagesize=A4)
            page_w, page_h = A4
            margin_x = 45
            margin_top = 55
            margin_bottom = 50
            line_h = 22
            y = page_h - margin_top
            page_num = 1

            def _draw_footer():
                c.setFont("Helvetica", 9)
                c.drawString(margin_x, 30, f"Page {page_num}")
                c.drawRightString(page_w - margin_x, 30, "Prometheus Document Engine")

            for raw_line in content.splitlines():
                line = raw_line.strip()
                if not line:
                    y -= line_h * 0.7
                    if y < margin_bottom + line_h:
                        _draw_footer()
                        c.showPage()
                        page_num += 1
                        y = page_h - margin_top
                    continue

                is_heading = line.startswith("#") or (line.startswith("*") and line.endswith("*") and len(line) < 60)
                clean_text = re.sub(r"^[#*]+s*", "", line).rstrip("*").strip()

                if is_heading:
                    font_size = 14
                    font_name = "Vazirmatn" if has_vazir else "Helvetica-Bold"
                    c.setFont(font_name, font_size)
                    y -= 6
                else:
                    font_size = 10
                    font_name = "Vazirmatn" if has_vazir else "Helvetica"
                    c.setFont(font_name, font_size)

                wrapped = textwrap.wrap(clean_text, width=65 if is_heading else 75)
                for w_line in wrapped:
                    if y < margin_bottom + line_h:
                        _draw_footer()
                        c.showPage()
                        page_num += 1
                        y = page_h - margin_top
                        c.setFont(font_name, font_size)

                    if _is_rtl(w_line):
                        shaped = _shape_persian(w_line)
                        c.drawRightString(page_w - margin_x, y, shaped)
                    else:
                        c.drawString(margin_x, y, w_line)
                    y -= line_h

            _draw_footer()
            c.showPage()
            c.save()

        elif ext == ".docx":
            # Generate Word Document
            import docx
            doc = docx.Document()
            doc.add_heading(base_name.replace("_", " ").title(), 0)
            for p in content.split("\n\n"):
                if p.strip():
                    doc.add_paragraph(p.strip())
            doc.save(file_bytes)

        elif ext == ".xlsx":
            # Generate Excel Spreadsheet
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            lines = content.strip().splitlines()
            for row_idx, line in enumerate(lines, 1):
                # Check comma or pipe or tab separated
                cells = [c.strip() for c in (line.split(",") if "," in line else (line.split("|") if "|" in line else line.split("\t")))]
                for col_idx, cell in enumerate(cells, 1):
                    ws.cell(row=row_idx, column=col_idx, value=cell)
            wb.save(file_bytes)

        elif ext == ".zip":
            import zipfile
            with zipfile.ZipFile(file_bytes, "w", zipfile.ZIP_DEFLATED) as zf:
                zf.writestr(f"{base_name}.txt", content.encode("utf-8"))

        elif ext == ".csv":
            # Generate CSV
            csv_str = io.StringIO()
            writer = csv.writer(csv_str)
            for line in content.strip().splitlines():
                row = [c.strip() for c in (line.split(",") if "," in line else (line.split("|") if "|" in line else [line]))]
                writer.writerow(row)
            file_bytes.write(csv_str.getvalue().encode("utf-8-sig"))

        elif ext == ".json":
            # Generate formatted JSON
            try:
                parsed_json = json.loads(content)
                file_bytes.write(json.dumps(parsed_json, indent=2, ensure_ascii=False).encode("utf-8"))
            except Exception:
                file_bytes.write(content.encode("utf-8"))

        else:
            # Plain Text / Code (PY, JS, HTML, CSS, MD, SH, SQL, etc.)
            file_bytes.write(content.encode("utf-8"))

        file_bytes.seek(0)
        raw_data = file_bytes.getvalue()

        return {
            "type": "document",
            "filename": clean_name,
            "bytes": raw_data,
            "caption": caption or f"📁 فایل `{clean_name}` با موفقیت تولید و آماده گردید.",
            "status": "success",
            "size_bytes": len(raw_data)
        }

    except Exception as e:
        logger.error(f"Error creating file {clean_name}: {e}")
        return {
            "type": "error",
            "status": "error",
            "message": f"خطا در تولید فایل {clean_name}: {str(e)}"
        }
